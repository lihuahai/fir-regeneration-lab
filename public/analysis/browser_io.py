import base64, binascii, csv, hashlib, io, json, math, re, time, zipfile
from pathlib import Path
import openpyxl
from engine import parse_table
MAX_ROWS, MAX_COLUMNS, MAX_BYTES = 300, 40, 5 * 1024 * 1024

def validate_dataset(dataset):
    if not isinstance(dataset, dict):
        raise ValueError('缺少有效数据表。')
    rows, columns = dataset.get('rows'), dataset.get('columns')
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise ValueError(f'在线版支持 1–{MAX_ROWS} 个样方。')
    if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_COLUMNS:
        raise ValueError(f'在线版最多支持 {MAX_COLUMNS} 列。')
    keys = set()
    for c in columns:
        if not isinstance(c, dict) or not all(isinstance(c.get(k), str) for k in ('key', 'label', 'role', 'unit')):
            raise ValueError('变量信息不完整，请重新导入文件。')
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,79}', c['key']) or c['key'] in keys:
            raise ValueError('变量编号不合法或重复。')
        if c['role'] not in ('id', 'label', 'count', 'trait', 'environment'):
            raise ValueError('变量用途不合法。')
        if max(len(c['label']), len(c['unit'])) > 200:
            raise ValueError('变量名称或单位过长。')
        keys.add(c['key'])
    if 'plot' not in keys:
        raise ValueError('需要保留“样方编号”列。')
    for row in rows:
        if not isinstance(row, dict) or len(row) > MAX_COLUMNS:
            raise ValueError('数据行格式不正确。')
        for key, value in row.items():
            if key not in keys:
                raise ValueError('数据行包含未定义的变量。')
            if value is not None and (isinstance(value, bool) or not isinstance(value, (str, int, float))):
                raise ValueError('单元格只支持文本、数字或空值。')
            if isinstance(value, str) and len(value) > 500:
                raise ValueError('单元格内容过长，请上传样方汇总数据。')
            if isinstance(value, (int, float)) and (not math.isfinite(value) or abs(value) > 1e12):
                raise ValueError('数字超出可分析范围。')
    # Do not accept client-supplied file paths or arbitrary metadata.
    return dict(name=str(dataset.get('name', '工作数据'))[:200],
                sheet=str(dataset.get('sheet', '数据'))[:100], columns=columns, rows=rows,
                layout=str(dataset.get('layout', '一行一个样方'))[:200],
                importIssues=[str(s)[:1000] for s in dataset.get('importIssues', [])[:50]])

def validate_config(config):
    if not isinstance(config, dict):
        raise ValueError('分析设置格式不正确。')
    result = {k: config[k] for k in ('target', 'low', 'high', 'predictors', 'quadratic',
                                   'cluster', 'family', 'area', 'areaColumn', 'roles') if k in config}
    for k in ('target', 'cluster', 'family', 'areaColumn'):
        if k in result and (not isinstance(result[k], str) or len(result[k]) > 80):
            raise ValueError('分析设置包含无效选项。')
    for k in ('low', 'high', 'area'):
        if k in result and result[k] not in ('', None):
            value = float(result[k])
            if not math.isfinite(value) or abs(value) > 1e12:
                raise ValueError('分析设置包含无效数值。')
    predictors = result.get('predictors', [])
    if not isinstance(predictors, list) or len(predictors) > 3 or any(not isinstance(k, str) or len(k) > 80 for k in predictors):
        raise ValueError('模型最多同时选择 3 个环境因子。')
    roles = result.get('roles', {})
    if not isinstance(roles, dict) or len(roles) > MAX_COLUMNS or any(v not in ('environment', 'label', 'count', 'trait') for v in roles.values()):
        raise ValueError('变量用途设置不正确。')
    return result

def source(dataset, mode):
    digest = hashlib.sha256(json.dumps(dataset, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:16]
    return dict(dataset=dataset, version=f'{mode}:{digest}', mode=mode,
                modified=time.time(), sourcePath='', sourceName=dataset['name'], config={})

def read_upload(body):
    name = str(body.get('name', ''))[:200].replace('\\', '/').rsplit('/', 1)[-1]
    ext = Path(name).suffix.lower()
    if ext not in ('.xlsx', '.csv'):
        raise ValueError('请上传 .xlsx 或 .csv 文件。')
    try:
        raw = base64.b64decode(body.get('base64', ''), validate=True)
    except (ValueError, TypeError, binascii.Error):
        raise ValueError('文件编码无效，请重新选择文件。') from None
    if not 0 < len(raw) <= MAX_BYTES:
        raise ValueError('单个文件大小应在 0–5 MB 之间。')
    matrix = []
    if ext == '.csv':
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('gb18030')
        for row in csv.reader(io.StringIO(content)):
            matrix.append(row)
            if len(matrix) > 1000 or len(row) > 80:
                raise ValueError('原表最多支持 1,000 行、80 列；合并后最多 300 个样方、40 个变量。')
        sheet_name = 'CSV'
    else:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 32 * 1024 * 1024:
                raise ValueError('Excel 解压内容过大，请删除无关工作表与图片后重试。')
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
        try:
            sheet = book.worksheets[0]
            sheet_name = sheet.title
            if (sheet.max_column or 0) > 80 or (sheet.max_row or 0) > 1000:
                raise ValueError('原表最多支持 1,000 行、80 列，请删除无关的空白行列。')
            for row in sheet.iter_rows(max_row=1001, max_col=81, values_only=True):
                if row[80] is not None:
                    raise ValueError('原表列数超出限制。')
                matrix.append(list(row[:80]))
            if len(matrix) > 1000 and any(v is not None for v in matrix[1000]):
                raise ValueError('原表行数超出限制。')
            matrix = matrix[:1000]
            while matrix and all(v is None for v in matrix[-1]):
                matrix.pop()
        finally:
            book.close()
    return validate_dataset(parse_table(matrix, name, sheet_name))
